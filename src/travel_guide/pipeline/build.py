from __future__ import annotations

import html
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .csv_io import read_csv
from .markdown_html import markdown_to_html
from .meta import load_meta
from .pdf import html_files_to_pdf

NAVY = "17365D"
PALE = "F5F7F8"
AMBER = "FCE8B2"
RED = "F4CCCC"
GREEN = "E2F0D9"
GRID = "CBD5DC"


def _style_sheet(ws, widths: list[int] | None = None) -> None:
    thin = Side(style="thin", color=GRID)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.row_dimensions[1].height = 36

    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        max_lines = 1
        risk_value = ""
        if ws.title == "每日行程" and ws.max_column >= 13:
            risk_value = str(ws.cell(row_idx, 13).value or "")
        fill = None
        if risk_value == "高":
            fill = PatternFill("solid", fgColor=RED)
        elif risk_value == "中":
            fill = PatternFill("solid", fgColor=AMBER)
        elif risk_value == "低":
            fill = PatternFill("solid", fgColor=GREEN)
        elif row_idx % 2 == 0:
            fill = PatternFill("solid", fgColor=PALE)

        for cell in row:
            text = str(cell.value or "")
            max_lines = max(max_lines, text.count("\n") + max(1, len(text) // 24))
            cell.font = Font(name="Microsoft YaHei", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if fill:
                cell.fill = fill
        ws.row_dimensions[row_idx].height = min(150, max(42, max_lines * 14))

    ws.freeze_panes = "A2"
    if ws.max_row and ws.max_column:
        ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False

    if widths:
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width
    else:
        for index in range(1, ws.max_column + 1):
            longest = max(
                (len(str(ws.cell(row, index).value or "")) for row in range(1, ws.max_row + 1)),
                default=12,
            )
            ws.column_dimensions[get_column_letter(index)].width = min(55, max(12, longest * 1.2))


def _add_csv_sheet(wb: Workbook, data_dir: Path, name: str, filename: str, widths: list[int] | None = None) -> None:
    headers, rows = read_csv(data_dir / filename)
    if not headers:
        return
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    _style_sheet(ws, widths)


def _build_workbooks(trip_root: Path, meta: dict) -> list[Path]:
    data_dir = trip_root / "data"
    output_dir = trip_root / "output"
    output_dir.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)
    _add_csv_sheet(wb, data_dir, "每日行程", "daily_itinerary.csv", [8, 14, 18, 23, 20, 19, 24, 18, 14, 14, 48, 44, 12, 44])
    for sheet_name, filename, widths in [
        ("酒店候选", "hotels.csv", [14, 28, 10, 14, 30, 14, 16, 22, 40]),
        ("每日三餐", "meals.csv", [10, 14, 10, 12, 18, 28, 12, 22, 18]),
        ("景点门票与交通", "attractions.csv", [18, 12, 20, 20, 26, 16, 16, 28, 28, 40]),
        ("天气风险与决策", "weather_risk.csv", [12, 20, 36, 36, 28, 24]),
        ("预算", "budget.csv", [22, 14, 14, 14, 28, 36]),
        ("预订清单", "booking_checklist.csv", [8, 28, 18, 24, 10, 40]),
        ("装备清单", "gear.csv", [14, 24, 10, 28, 24]),
        ("应急预案", "contingencies.csv", [20, 28, 32, 32, 28]),
        ("参考来源", "sources.csv", [8, 14, 28, 50, 36, 24]),
    ]:
        _add_csv_sheet(wb, data_dir, sheet_name, filename, widths)

    itinerary_path = output_dir / meta["xlsx_itinerary"]
    wb.save(itinerary_path)

    nav_wb = Workbook()
    nav_wb.remove(nav_wb.active)
    for sheet_name, filename, widths in [
        ("摄影GPS点位", "waypoints.csv", [14, 28, 12, 12, 16, 18, 18, 22, 36]),
        ("景区入口与交通", "attractions.csv", [18, 12, 20, 20, 26, 16, 16, 28, 28, 40]),
        ("参考来源", "sources.csv", [8, 14, 28, 50, 36, 24]),
    ]:
        _add_csv_sheet(nav_wb, data_dir, sheet_name, filename, widths)
    nav_path = output_dir / meta["xlsx_waypoints"]
    nav_wb.save(nav_path)
    return [itinerary_path, nav_path]


def _csv_to_html(data_dir: Path, title: str, filename: str) -> str:
    headers, rows = read_csv(data_dir / filename)
    if not headers:
        return ""
    output = [f'<section class="appendix"><h1>{html.escape(title)}</h1><table><thead><tr>']
    output.extend(f"<th>{html.escape(header)}</th>" for header in headers)
    output.append("</tr></thead><tbody>")
    for row in rows:
        output.append("<tr>")
        for value in row:
            escaped = html.escape(value).replace("\n", "<br>")
            if value.startswith(("http://", "https://")):
                escaped = f'<a href="{html.escape(value)}">{html.escape(value)}</a>'
            output.append(f"<td>{escaped}</td>")
        output.append("</tr>")
    output.append("</tbody></table></section>")
    return "".join(output)


def _build_html(trip_root: Path, meta: dict) -> list[Path]:
    content_dir = trip_root / "content"
    data_dir = trip_root / "data"
    map_dir = trip_root / "maps"
    output_dir = trip_root / "output"
    output_dir.mkdir(parents=True, exist_ok=True)

    chapters = [markdown_to_html(path.read_text(encoding="utf-8")) for path in sorted(content_dir.glob("*.md"))]
    for title, filename in [
        ("附录A：住宿候选与预订核验", "hotels.csv"),
        ("附录B：每日三餐执行表", "meals.csv"),
        ("附录C：天气风险与动态决策", "weather_risk.csv"),
        ("附录D：景点门票与景区交通", "attractions.csv"),
        ("附录E：摄影GPS点位", "waypoints.csv"),
        ("附录F：预算明细", "budget.csv"),
        ("附录G：装备清单", "gear.csv"),
        ("附录H：应急预案", "contingencies.csv"),
        ("附录I：预订清单", "booking_checklist.csv"),
        ("附录J：参考来源", "sources.csv"),
    ]:
        chapters.append(_csv_to_html(data_dir, title, filename))

    overview = map_dir / "overview-route.svg"
    site = map_dir / "site-flows.svg"
    overview_uri = overview.resolve().as_uri() if overview.exists() else ""
    site_uri = site.resolve().as_uri() if site.exists() else ""
    facts = "".join(f"<p>{html.escape(item)}</p>" for item in meta.get("cover_facts") or [])

    map_blocks = ""
    if overview_uri:
        map_blocks += f'<section class="map-page"><h1>总路线图</h1><img src="{overview_uri}" alt="overview"></section>'
    if site_uri:
        map_blocks += f'<section class="map-page"><h1>核心景区交通链</h1><img src="{site_uri}" alt="site flows"></section>'

    document = f"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<title>{html.escape(meta["title"])}</title>
<style>
@page {{ size: A4; margin: 18mm 15mm 18mm; }}
@page wide {{ size: A4 landscape; margin: 12mm; }}
* {{ box-sizing: border-box; }}
body {{ margin: 0; color: #24313a; font-family: "Microsoft YaHei", "Noto Sans CJK SC", sans-serif; font-size: 10.5pt; line-height: 1.62; }}
.cover {{ min-height: 255mm; display: flex; flex-direction: column; justify-content: center; border-top: 8px solid #17365d; page-break-after: always; }}
.cover h1 {{ font-size: 28pt; line-height: 1.2; margin: 0 0 16px; color: #17365d; }}
.cover .subtitle {{ font-size: 16pt; color: #53606d; margin-bottom: 50px; }}
.cover .facts {{ border-left: 4px solid #c55a11; padding-left: 20px; font-size: 12pt; }}
.map-page {{ page-break-after: always; }}
.map-page img {{ width: 100%; margin-top: 20px; }}
h1 {{ color: #17365d; font-size: 24pt; margin-top: 0; page-break-before: always; }}
h2 {{ color: #17365d; font-size: 17pt; border-bottom: 2px solid #d9eaf7; padding-bottom: 6px; margin-top: 30px; }}
h3 {{ color: #315b78; font-size: 13pt; margin-top: 22px; }}
p {{ margin: 8px 0; }}
table {{ width: 100%; border-collapse: collapse; margin: 12px 0 20px; font-size: 8.5pt; }}
th {{ background: #17365d; color: white; text-align: left; padding: 7px; border: 1px solid #cbd5dc; }}
td {{ vertical-align: top; padding: 7px; border: 1px solid #cbd5dc; }}
tr:nth-child(even) td {{ background: #f5f7f8; }}
a {{ color: #1e5b87; overflow-wrap: anywhere; }}
.footer-note {{ margin-top: 30px; font-size: 8.5pt; color: #6d747b; }}
.appendix {{ page: wide; page-break-before: always; }}
.appendix h1 {{ page-break-before: auto; }}
.appendix table {{ font-size: 7pt; line-height: 1.35; table-layout: fixed; }}
.appendix th, .appendix td {{ padding: 4px; overflow-wrap: anywhere; }}
</style>
</head>
<body>
<section class="cover">
  <h1>{html.escape(meta["title"])}</h1>
  <div class="subtitle">{html.escape(meta.get("subtitle") or "")}</div>
  <div class="facts">{facts}</div>
  <p class="footer-note">{html.escape(meta.get("footer_note") or "")}</p>
</section>
{map_blocks}
{"".join(chapters)}
</body>
</html>"""
    handbook = output_dir / meta["html_handbook"]
    handbook.write_text(document, encoding="utf-8")

    route_parts = []
    if overview_uri:
        route_parts.append(f'<section class="page"><img src="{overview_uri}" alt="overview"></section>')
    if site_uri:
        route_parts.append(f'<section class="page"><img src="{site_uri}" alt="site"></section>')
    if not route_parts:
        route_parts.append('<section class="page"><p>暂无地图 SVG，请先生成 maps/overview-route.svg</p></section>')

    route_html = f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><title>{html.escape(meta["title"])} A3</title>
<style>
@page {{ size: A3 landscape; margin: 10mm; }}
body {{ margin: 0; font-family: "Microsoft YaHei", Arial, sans-serif; color: #17365d; }}
.page {{ height: 277mm; display: flex; align-items: center; justify-content: center; page-break-after: always; }}
.page:last-child {{ page-break-after: auto; }}
img {{ width: 100%; max-height: 270mm; object-fit: contain; }}
</style></head><body>{"".join(route_parts)}</body></html>"""
    map_html = output_dir / meta["html_map"]
    map_html.write_text(route_html, encoding="utf-8")
    return [handbook, map_html]


def ensure_placeholder_maps(trip_root: Path, meta: dict) -> None:
    """Create simple SVG placeholders if maps are missing."""
    map_dir = trip_root / "maps"
    map_dir.mkdir(parents=True, exist_ok=True)
    overview = map_dir / "overview-route.svg"
    site = map_dir / "site-flows.svg"
    title = html.escape(meta.get("title") or "Trip")
    if not overview.exists():
        overview.write_text(
            f"""<?xml version="1.0" encoding="UTF-8"?>
<svg xmlns="http://www.w3.org/2000/svg" width="1600" height="900" viewBox="0 0 1600 900">
  <rect width="1600" height="900" fill="#f7f4ed"/>
  <text x="80" y="120" font-size="40" font-family="Microsoft YaHei,sans-serif" fill="#17365d">{title}</text>
  <text x="80" y="180" font-size="22" fill="#53606d">Overview map placeholder — replace with generate_maps output</text>
  <circle cx="300" cy="450" r="14" fill="#17365d"/><circle cx="700" cy="380" r="14" fill="#17365d"/>
  <circle cx="1100" cy="500" r="14" fill="#17365d"/><circle cx="1400" cy="420" r="14" fill="#17365d"/>
  <polyline points="300,450 700,380 1100,500 1400,420" fill="none" stroke="#d56a22" stroke-width="8"/>
</svg>
""",
            encoding="utf-8",
        )
    if not site.exists():
        site.write_text(
            f"""<?xml version="1.0" encoding="UTF-8"?>
<svg xmlns="http://www.w3.org/2000/svg" width="1600" height="900" viewBox="0 0 1600 900">
  <rect width="1600" height="900" fill="#f4f7f2"/>
  <text x="80" y="120" font-size="36" font-family="Microsoft YaHei,sans-serif" fill="#17365d">Site flow placeholder</text>
  <text x="80" y="180" font-size="20" fill="#53606d">{title}</text>
  <line x1="120" y1="450" x2="1480" y2="450" stroke="#2274a5" stroke-width="6"/>
  <circle cx="120" cy="450" r="12" fill="#17365d"/><circle cx="560" cy="450" r="12" fill="#17365d"/>
  <circle cx="1000" cy="450" r="12" fill="#17365d"/><circle cx="1480" cy="450" r="12" fill="#17365d"/>
</svg>
""",
            encoding="utf-8",
        )


def build_trip(trip_root: Path, *, make_pdf: bool = True) -> list[Path]:
    trip_root = trip_root.resolve()
    meta = load_meta(trip_root)
    ensure_placeholder_maps(trip_root, meta)
    paths = []
    paths.extend(_build_workbooks(trip_root, meta))
    html_paths = _build_html(trip_root, meta)
    paths.extend(html_paths)
    if make_pdf:
        paths.extend(html_files_to_pdf(trip_root, meta))
    return paths
