from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

from .meta import load_meta
from .csv_io import read_csv


def _verify_xlsx(path: Path, required_sheets: list[str]) -> None:
    if not path.exists() or path.stat().st_size < 5_000:
        raise RuntimeError(f"XLSX missing or too small: {path}")
    with ZipFile(path) as archive:
        if "[Content_Types].xml" not in archive.namelist():
            raise RuntimeError(f"Not a valid Office Open XML package: {path}")
    workbook = load_workbook(path, read_only=True, data_only=False)
    missing = [sheet for sheet in required_sheets if sheet not in workbook.sheetnames]
    if missing:
        raise RuntimeError(f"Missing sheets in {path.name}: {missing}")
    for sheet in required_sheets:
        ws = workbook[sheet]
        if ws.max_row < 2 or ws.max_column < 2:
            raise RuntimeError(f"Empty sheet {sheet} in {path.name}")
    workbook.close()


def _verify_html(path: Path, required: list[str], forbidden: list[str]) -> None:
    if not path.exists() or path.stat().st_size < 5_000:
        raise RuntimeError(f"HTML missing or too small: {path}")
    text = path.read_text(encoding="utf-8")
    missing = [term for term in required if term and term not in text]
    if missing:
        raise RuntimeError(f"Missing terms in {path.name}: {missing}")
    stale = [term for term in forbidden if term and term in text]
    if stale:
        raise RuntimeError(f"Stale terms in {path.name}: {stale}")


def _verify_pdf(path: Path) -> None:
    if not path.exists() or path.stat().st_size < 20_000:
        raise RuntimeError(f"PDF missing or too small: {path}")
    if path.read_bytes()[:5] != b"%PDF-":
        raise RuntimeError(f"Invalid PDF header: {path}")


def verify_budget(trip_root: Path) -> None:
    headers, rows = read_csv(trip_root / "data" / "budget.csv")
    if not rows:
        return
    # Find 小计 row; sum numeric rows before it
    subtotal_idx = next((i for i, row in enumerate(rows) if row and row[0] == "小计"), None)
    if subtotal_idx is None:
        return
    item_rows = rows[:subtotal_idx]
    for col in (1, 2, 3):
        total = 0
        for row in item_rows:
            if len(row) <= col or not str(row[col]).strip():
                continue
            try:
                total += int(float(row[col]))
            except ValueError:
                continue
        stated = int(float(rows[subtotal_idx][col]))
        if total != stated:
            raise RuntimeError(f"budget.csv column {col + 1} sum {total} != 小计 {stated}")


def verify_trip(trip_root: Path, *, require_pdf: bool = True) -> None:
    trip_root = trip_root.resolve()
    meta = load_meta(trip_root)
    output = trip_root / "output"
    verify_budget(trip_root)
    _verify_xlsx(output / meta["xlsx_itinerary"], ["每日行程", "预算"])
    _verify_xlsx(output / meta["xlsx_waypoints"], ["摄影GPS点位"])
    required = list(meta.get("required_terms") or [])
    if meta.get("title"):
        required = [meta["title"], *required]
    _verify_html(output / meta["html_handbook"], required, list(meta.get("forbidden_terms") or []))
    if require_pdf:
        _verify_pdf(output / meta["pdf_handbook"])
        map_pdf = output / meta["pdf_map"]
        if map_pdf.exists():
            _verify_pdf(map_pdf)
